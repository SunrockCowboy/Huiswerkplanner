from datetime import date, timedelta
import io
import sys
import streamlit as st

# Word-document bibliotheek
HAS_DOCX = True
try:
  from docx import Document
  from docx.shared import Pt, RGBColor
except Exception:
  HAS_DOCX = False

st.set_page_config(
    page_title="Huiswerkplanner voor bazen",
    page_icon="📚",
    layout="centered",
)

st.title("📚 Huiswerkplanner voor bazen")
st.write(
    "Kies hieronder voor hoeveel toetsen je een planning wilt maken, vul per"
    " vak de onderdelen in en vink af wat klaar is!"
)

# --- INITIALISATIE VAN DE SESSIE STATUS ---
if "planning_gegenereerd" not in st.session_state:
  st.session_state.planning_gegenereerd = False

if "dagen_planning" not in st.session_state:
  st.session_state.dagen_planning = []

if "foutmelding" not in st.session_state:
  st.session_state.foutmelding = ""

if "toetsen" not in st.session_state:
  # Standaard starten we nu met 1 toets
  st.session_state.toetsen = [
      {
          "vak": "Duits",
          "toetsdatum": date.today() + timedelta(days=2),
          "leertaken": [
              {
                  "eenheid": "woordjes",
                  "aantal": None,
                  "van_blz": 1,
                  "tot_blz": 10,
                  "van_par": "3.1",
                  "tot_par": "3.4",
                  "wat_omschrijving": "",
                  "gekozen_dag_str": "🔄 Elke dag",
              }
          ],
      }
  ]


# --- FUNCTIES ---
def voeg_taak_toe_aan_toets(toets_idx):
  st.session_state.toetsen[toets_idx]["leertaken"].append({
      "eenheid": "woordjes",
      "aantal": None,
      "van_blz": 1,
      "tot_blz": 5,
      "van_par": "1.1",
      "tot_par": "1.2",
      "wat_omschrijving": "",
      "gekozen_dag_str": "🔄 Elke dag",
  })


def verwijder_taak_van_toets(toets_idx, taak_idx):
  if len(st.session_state.toetsen[toets_idx]["leertaken"]) > 1:
    st.session_state.toetsen[toets_idx]["leertaken"].pop(taak_idx)


def verwijder_toets(toets_idx):
  if len(st.session_state.toetsen) > 1:
    st.session_state.toetsen.pop(toets_idx)
    if "aantal_toetsen_selectie" in st.session_state:
      del st.session_state["aantal_toetsen_selectie"]

    keys_te_wissen = [
        k
        for k in st.session_state.keys()
        if k.startswith(
            (
                "vak_",
                "toetsdatum_",
                "eenheid_",
                "van_",
                "tot_",
                "aantal_",
                "dag_select_",
                "wat_",
            )
        )
    ]
    for k in keys_te_wissen:
      if k in st.session_state:
        del st.session_state[k]


def genereer_planning():
  vandaag = date.today()
  st.session_state.foutmelding = ""
  st.session_state.planning_gegenereerd = True

  alle_toetsdatums = []
  for t_idx, toets in enumerate(st.session_state.toetsen):
    t_datum = st.session_state.get(f"toetsdatum_{t_idx}", toets["toetsdatum"])
    alle_toetsdatums.append(t_datum)
    if (t_datum - vandaag).days <= 0:
      st.session_state.planning_gegenereerd = False
      st.session_state.foutmelding = (
          "⚠️ Kies voor alle toetsen een datum in de toekomst (op z'n"
          " vroegst morgen)!"
      )
      return

  if not st.session_state.planning_gegenereerd:
    return

  laatste_toetsdatum = max(alle_toetsdatums)

  beschikbare_datums = []
  d = vandaag + timedelta(days=1)
  while d <= laatste_toetsdatum:
    beschikbare_datums.append(d)
    d += timedelta(days=1)

  if not beschikbare_datums:
    beschikbare_datums = [vandaag + timedelta(days=1)]

  geconstrueerde_dagen = {}
  for dag_d in beschikbare_datums:
    geconstrueerde_dagen[dag_d] = {
        "header": f"🗓️ {dag_d.strftime('%d-%m-%Y')}",
        "items": [],
    }

  for t_idx, toets in enumerate(st.session_state.toetsen):
    vak_naam = st.session_state.get(f"vak_{t_idx}", toets["vak"])
    t_datum = st.session_state.get(f"toetsdatum_{t_idx}", toets["toetsdatum"])

    toets_specifieke_datums = []
    d_sub = vandaag + timedelta(days=1)
    while d_sub < t_datum:
      toets_specifieke_datums.append(d_sub)
      d_sub += timedelta(days=1)

    if not toets_specifieke_datums:
      toets_specifieke_datums = [vandaag + timedelta(days=1)]

    for taak_idx, taak in enumerate(toets["leertaken"]):
      eenheid = taak["eenheid"]
      omschrijving = (
          f" ({taak['wat_omschrijving']})" if taak["wat_omschrijving"] else ""
      )
      gekozen_dag_str = st.session_state.get(
          f"dag_select_{t_idx}_{taak_idx}",
          taak.get("gekozen_dag_str", "🔄 Elke dag"),
      )

      if eenheid == "bladzijden":
        van_nr = taak.get("van_blz", 1)
        tot_nr = taak.get("tot_blz", 1)
        if van_nr == tot_nr:
          tekst = f"[{vak_naam}] 📖 Lees blz. {van_nr}{omschrijving}"
        else:
          tekst = (
              f"[{vak_naam}] 📖 Lees blz. {van_nr} t/m {tot_nr}{omschrijving}"
          )

      elif eenheid == "paragrafen":
        van_par = taak.get("van_par", "1")
        tot_par = taak.get("tot_par", "1")
        tekst = (
            f"[{vak_naam}] 📑 Paragraaf {van_par} t/m"
            f" {tot_par}{omschrijving}"
        )

      else:
        aantal = taak.get("aantal", 10)
        aantal_str = str(aantal) if aantal is not None else "?"
        tekst = f"[{vak_naam}] 🔤 Leer {aantal_str} {eenheid}{omschrijving}"

      if gekozen_dag_str == "🔄 Elke dag":
        for dag_d in toets_specifieke_datums:
          if dag_d in geconstrueerde_dagen:
            geconstrueerde_dagen[dag_d]["items"].append(tekst)
      else:
        doel_datum = toets_specifieke_datums[0]
        for bd in toets_specifieke_datums:
          if bd.strftime("%d-%m-%Y") == gekozen_dag_str:
            doel_datum = bd
            break
        if doel_datum in geconstrueerde_dagen:
          geconstrueerde_dagen[doel_datum]["items"].append(tekst)

  gefilterde_dagen = []
  for dag_d, info in geconstrueerde_dagen.items():
    if info["items"]:
      gefilterde_dagen.append(info)

  st.session_state.dagen_planning = gefilterde_dagen


def maak_word_document(planning_data):
  doc = Document()
  p_title = doc.add_paragraph()
  run_title = p_title.add_run("📚 Gecombineerde Studieplanning")
  run_title.bold = True
  run_title.font.size = Pt(20)

  doc.add_paragraph(f"Gemaakt op: {date.today().strftime('%d-%m-%Y')}")
  doc.add_paragraph("--------------------------------------------------")

  for dag in planning_data:
    dag_titel = dag.get("header", "Dag")
    items = dag.get("items", [])

    if items:
      p_dag = doc.add_paragraph()
      run_dag = p_dag.add_run(dag_titel)
      run_dag.bold = True
      run_dag.font.size = Pt(14)

      for item in items:
        doc.add_paragraph(f"[  ]  {item}")

      doc.add_paragraph("")

  stream = io.BytesIO()
  doc.save(stream)
  stream.seek(0)
  return stream


# --- SELECTIE VOOR AANTAL TOETSEN (1 T/M 10) ---
st.header("1. Toetsen Inplannen")

huidig_aantal = len(st.session_state.toetsen)
aantal_toetsen = st.selectbox(
    "Voor hoeveel toetsen wil je een planning invullen?",
    options=list(range(1, 11)),
    index=huidig_aantal - 1 if 1 <= huidig_aantal <= 10 else 0,
    key="aantal_toetsen_selectie",
)

if aantal_toetsen != len(st.session_state.toetsen):
  if aantal_toetsen > len(st.session_state.toetsen):
    standaard_vakken = [
        "Wiskunde",
        "Engels",
        "Frans",
        "Aardrijkskunde",
        "Natuurkunde",
        "Scheikunde",
        "Economie",
        "Maatschappijleer",
    ]
    while len(st.session_state.toetsen) < aantal_toetsen:
      idx = len(st.session_state.toetsen)
      vak_naam = standaard_vakken[idx % len(standaard_vakken)]
      st.session_state.toetsen.append({
          "vak": vak_naam,
          "toetsdatum": date.today() + timedelta(days=(idx + 1) * 2),
          "leertaken": [
              {
                  "eenheid": "woordjes",
                  "aantal": None,
                  "van_blz": 1,
                  "tot_blz": 10,
                  "van_par": "1.1",
                  "tot_par": "1.2",
                  "wat_omschrijving": "",
                  "gekozen_dag_str": "🔄 Elke dag",
              }
          ],
      })
  else:
    st.session_state.toetsen = st.session_state.toetsen[:aantal_toetsen]
  st.rerun()

st.markdown("---")

# --- DYNAMISCH INVOERVELD PER TOETS ---
for t_idx, toets in enumerate(st.session_state.toetsen):
  with st.container():
    c_titel, c_del_toets = st.columns([3, 1])
    with c_titel:
      st.subheader(f"📌 Toets {t_idx+1}")
    with c_del_toets:
      if len(st.session_state.toetsen) > 1:
        if st.button("🗑️ Wis toets", key=f"del_toets_{t_idx}"):
          verwijder_toets(t_idx)
          st.rerun()

    col_v, col_d = st.columns(2)
    with col_v:
      vak_val = st.text_input(
          f"Schoolvak {t_idx+1}",
          value=toets["vak"],
          key=f"vak_{t_idx}",
      )
      st.session_state.toetsen[t_idx]["vak"] = vak_val
    with col_d:
      morgen = date.today() + timedelta(days=1)
      toets_val = st.date_input(
          f"Toetsdatum vak {t_idx+1}",
          value=toets.get("toetsdatum", morgen),
          key=f"toetsdatum_{t_idx}",
      )
      st.session_state.toetsen[t_idx]["toetsdatum"] = toets_val

    huidige_t_datum = st.session_state.get(f"toetsdatum_{t_idx}", morgen)
    beschikbare_opties = ["🔄 Elke dag"]
    d_loop = date.today() + timedelta(days=1)
    while d_loop < huidige_t_datum:
      beschikbare_opties.append(d_loop.strftime("%d-%m-%Y"))
      d_loop += timedelta(days=1)

    st.markdown(f"*Wat moet je voor {vak_val} leren?*")

    for taak_idx, taak in enumerate(toets["leertaken"]):
      c_taak_titel, c_del_taak = st.columns([3, 1])
      with c_taak_titel:
        st.markdown(f"&nbsp;&nbsp;&nbsp;&bull; *Onderdeel {taak_idx+1}*")
      with c_del_taak:
        if len(toets["leertaken"]) > 1:
          if st.button("🗑️ Wis", key=f"del_taak_{t_idx}_{taak_idx}"):
            verwijder_taak_van_toets(t_idx, taak_idx)
            st.rerun()

      col1, col2, col3 = st.columns([2, 2, 2])

      with col1:
        gekozen_eenheid = st.selectbox(
            "Soort stof",
            [
                "woordjes",
                "bladzijden",
                "paragrafen",
                "sommen",
                "grammatica-regels",
            ],
            index=[
                "woordjes",
                "bladzijden",
                "paragrafen",
                "sommen",
                "grammatica-regels",
            ].index(taak["eenheid"])
            if taak["eenheid"]
            in [
                "woordjes",
                "bladzijden",
                "paragrafen",
                "sommen",
                "grammatica-regels",
            ]
            else 0,
            key=f"eenheid_{t_idx}_{taak_idx}",
        )
        toets["leertaken"][taak_idx]["eenheid"] = gekozen_eenheid

      with col2:
        if gekozen_eenheid == "bladzijden":
          subcol1, subcol2 = st.columns(2)
          with subcol1:
            toets["leertaken"][taak_idx]["van_blz"] = st.number_input(
                "Van blz.",
                min_value=1,
                value=taak.get("van_blz", 1),
                key=f"van_{t_idx}_{taak_idx}",
            )
          with subcol2:
            toets["leertaken"][taak_idx]["tot_blz"] = st.number_input(
                "T/m blz.",
                min_value=1,
                value=taak.get("tot_blz", 10),
                key=f"tot_{t_idx}_{taak_idx}",
            )

        elif gekozen_eenheid == "paragrafen":
          subcol1, subcol2 = st.columns(2)
          with subcol1:
            toets["leertaken"][taak_idx]["van_par"] = st.text_input(
                "Van par.",
                value=str(taak.get("van_par", "3.1")),
                key=f"van_p_{t_idx}_{taak_idx}",
            )
          with subcol2:
            toets["leertaken"][taak_idx]["tot_par"] = st.text_input(
                "T/m par.",
                value=str(taak.get("tot_par", "3.4")),
                key=f"tot_p_{t_idx}_{taak_idx}",
            )

        else:
          toets["leertaken"][taak_idx]["aantal"] = st.number_input(
              "Aantal",
              min_value=1,
              value=taak.get("aantal", None),
              key=f"aantal_{t_idx}_{taak_idx}",
          )

      with col3:
        huidige_val = taak.get("gekozen_dag_str", "🔄 Elke dag")
        if huidige_val not in beschikbare_opties:
          huidige_val = beschikbare_opties[0]

        gekozen_dag_str = st.selectbox(
            "Op welke dag leren?",
            beschikbare_opties,
            index=(
                beschikbare_opties.index(huidige_val)
                if huidige_val in beschikbare_opties
                else 0
            ),
            key=f"dag_select_{t_idx}_{taak_idx}",
        )
        toets["leertaken"][taak_idx]["gekozen_dag_str"] = gekozen_dag_str

      toets["leertaken"][taak_idx]["wat_omschrijving"] = st.text_input(
          "Wat precies? (bijv. 'Alleen werkwoorden')",
          value=taak.get("wat_omschrijving", ""),
          key=f"wat_{t_idx}_{taak_idx}",
      )
      st.markdown("")

    if st.button(
        f"➕ Voeg extra onderdeel toe voor {vak_val}", key=f"add_taak_{t_idx}"
    ):
      voeg_taak_toe_aan_toets(t_idx)
      st.rerun()

    st.divider()

st.write("")
st.button(
    "🚀 Genereer Planning voor alle geselecteerde toetsen",
    on_click=genereer_planning,
)

if st.session_state.foutmelding:
  st.error(st.session_state.foutmelding)

# --- OVERZICHT, DOWNLOAD & AFVINKEN ---
if st.session_state.get("planning_gegenereerd", False):
  st.divider()
  st.subheader("📌 Jouw Gecombineerde Studieplanning:")

  huidige_planning = st.session_state.get("dagen_planning", [])

  if HAS_DOCX:
    word_stream = maak_word_document(huidige_planning)
    st.download_button(
        label="📄 Download als Word-document (.docx)",
        data=word_stream,
        file_name=f"Studieplanning_{date.today().strftime('%d-%m-%Y')}.docx",
        mime=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
  else:
    st.warning("Installeer python-docx voor Word-export.")

  st.markdown("---")
  st.subheader("✅ Vink hier je geleerde taken af:")

  heeft_taken = False
  for dag_idx, dag in enumerate(huidige_planning):
    dag_titel = dag.get("header", f"Dag {dag_idx+1}")
    items = dag.get("items", [])

    if items:
      heeft_taken = True
      st.markdown(f"**{dag_titel}**")
      for item_idx, item_tekst in enumerate(items):
        checkbox_key = f"check_{dag_idx}_{item_idx}_{item_tekst}"
        is_gelukt = st.checkbox(item_tekst, key=checkbox_key)

        if is_gelukt:
          st.success(f"🎉 **Top gedaan!** ({item_tekst})")
      st.write("")

  if not heeft_taken:
    st.info("Er zijn nog geen taken ingepland.")
